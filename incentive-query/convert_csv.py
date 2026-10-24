# -*- coding: utf-8 -*-
"""把三 Sheet 的 xlsx 明细转换为三个同名 CSV（服务端依赖装不上时的保底）。
用法：python convert_csv.py [输入.xlsx]
默认：data/incentive_sample.xlsx -> data/试驾激励.csv / 交付激励.csv / 终端激励.csv
"""
import sys
import os
import csv
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "data", "incentive_sample.xlsx")
out_dir = os.path.dirname(src)

wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    rows = [r for r in ws.iter_rows(values_only=True)]
    dst = os.path.join(out_dir, f"{name}.csv")
    with open(dst, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for r in rows:
            w.writerow(["" if c is None else c for c in r])
    print("converted:", name, "->", dst, "rows:", len(rows) - 1)
wb.close()
