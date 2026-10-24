# -*- coding: utf-8 -*-
"""生成数据模板：data/激励明细模板.xlsx
Sheet1 填写说明（含三张原表的对应关系与查询钥匙说明）
Sheet2-4 试驾激励 / 交付激励 / 终端激励（列 = 原表列原序 + 末列「月份」，各带1行示例）
三张原表可整列直接粘贴，只需在末尾补「月份」列。
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "data", "激励明细模板.xlsx")

COLS_SY = ["大区", "经销商编码", "经销商名称", "门店类型", "车系", "上牌日期",
           "销售名字", "主岗位", "在职", "身份证", "激励金额", "月份"]
COLS_JF = ["大区", "门店类型", "经销商代码", "经销商名称", "员工工号", "姓名",
           "身份证", "APP订单号", "DMS车系名称", "车型标准名称", "激励金额", "月份"]
COLS_ZD = ["大区", "序号", "店编号", "店简称", "店类型", "总经理名称", "总经理电话",
           "总经理身份证", "N60终端目标", "N60实绩", "激励金额", "月份"]

wb = openpyxl.Workbook()

# ============ Sheet1 填写说明 ============
ws = wb.active
ws.title = "填写说明"
title_font = Font(bold=True, size=14)
h_font = Font(bold=True, color="FFFFFF")
h_fill = PatternFill("solid", fgColor="185FA5")
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(*[Side(style="thin", color="D0D5DD")] * 4)

ws["A1"] = "激励明细模板 · 填写说明（列结构与三张原表完全一致，可直接整列粘贴）"
ws["A1"].font = title_font

notes = [
    ["一、结构", "三个 Sheet（试驾激励/交付激励/终端激励）分别对应你们的三张原表，列名与原表完全一致，只在末尾多了一列「月份」——这是查询用的辅助列，请填 2026-09 这种「4位年-2位月」格式。"],
    ["二、怎么填", "把当月三张原表的内容整列粘贴进对应 Sheet，最后补上月份列即可。序号、总经理电话等列原表有就带着，不影响查询。"],
    ["三、查询钥匙", "店端输入「店代码」：自动匹配 经销商编码(试驾) / 经销商代码(交付) / 店编号(终端)，一次查回本店三类激励；个人输入「完整身份证号」：自动匹配 身份证(试驾/交付) + 总经理身份证(终端)，一次查回本人全部激励。系统自动识别输入的是店代码还是身份证，无需选择。"],
    ["四、留痕", "每次查询自动记录：谁查的（店代码/掩码身份证+查询人姓名）、什么时候、查了几次、是否查到、合计金额，写入 data/query_log.csv，可随时审计。查询页无导出/打印功能。"],
    ["五、更新数据", "填好的表覆盖 data/incentive_sample.xlsx（Sheet 名保持不变），重新发布即可生效，无需改代码。"],
]
r = 3
for a, b in notes:
    ws.cell(row=r, column=1, value=a).font = Font(bold=True)
    c = ws.cell(row=r, column=2, value=b)
    c.alignment = wrap
    ws.row_dimensions[r].height = 46
    r += 1

r += 1
ws.cell(row=r, column=1, value="六、原表列 → 查询匹配逻辑对照").font = Font(bold=True, size=12)
r += 1
map_header = ["Sheet", "店代码查询匹配列", "身份证查询匹配列"]
for i, h in enumerate(map_header, start=1):
    c = ws.cell(row=r, column=i, value=h)
    c.font = h_font
    c.fill = h_fill
    c.border = thin
mapping = [
    ["试驾激励", "经销商编码", "身份证"],
    ["交付激励", "经销商代码", "身份证"],
    ["终端激励", "店编号", "总经理身份证"],
]
for m in mapping:
    r += 1
    for i, v in enumerate(m, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = thin
        c.alignment = wrap

ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 60
ws.column_dimensions["C"].width = 24

# ============ Sheet2-4 三类激励 ============
yellow = PatternFill("solid", fgColor="FFF7CC")
examples = {
    "试驾激励": ["华南", "GZ001", "广州昊铂体验中心", "体验店", "昊铂GT", "2026-10-08", "张伟", "销售顾问", "是", "440101199001011234", 300, "2026-10"],
    "交付激励": ["华南", "体验店", "GZ001", "广州昊铂体验中心", "YG1001", "张伟", "440101199001011234", "APP2026100105", "昊铂GT", "昊铂GT 710后驱版", 2000, "2026-10"],
    "终端激励": ["华南", 1, "GZ001", "广州昊铂体验中心", "体验店", "陈刚", "13800000001", "440104198505055678", 10, 9, 4500, "2026-10"],
}
for name, cols in [("试驾激励", COLS_SY), ("交付激励", COLS_JF), ("终端激励", COLS_ZD)]:
    ws2 = wb.create_sheet(name)
    for i, col in enumerate(cols, start=1):
        c = ws2.cell(row=1, column=i, value=col)
        c.font = h_font
        c.fill = h_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws2.column_dimensions[get_column_letter(i)].width = max(10, len(col) * 2.2)
    ws2.append(examples[name])
    for cc in range(1, len(cols) + 1):
        ws2.cell(row=2, column=cc).fill = yellow
    ws2.cell(row=4, column=1, value="↑ 黄色1行为示例，正式填写时请删除。末列「月份」为查询辅助列，必填。").font = Font(color="92400E")
    ws2.freeze_panes = "A2"

wb.save(OUT)
print("template saved:", OUT)
